import pytest

pptx = pytest.importorskip("pptx")
from pptx.util import Emu

from extract_slides_text import extract, parse_pages, render


def _deck(tmp_path):
    prs = pptx.Presentation()
    blank = prs.slide_layouts[6]

    s1 = prs.slides.add_slide(blank)
    body = s1.shapes.add_textbox(Emu(100), Emu(900000), Emu(5000000), Emu(1000000))
    body.text_frame.text = "최적화를 위한 하이브리드 차량 축약 모델 구성"
    body.text_frame.add_paragraph().text = "1) 하이브리드 차량 축약 모델 구성"
    top = s1.shapes.add_textbox(Emu(100), Emu(180000), Emu(3000000), Emu(300000))
    top.text_frame.text = "주요 연구내용"
    table = s1.shapes.add_table(2, 2, Emu(100), Emu(3000000), Emu(3000000), Emu(600000)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "항목", "값"
    table.cell(1, 0).text, table.cell(1, 1).text = "SoC 상한", "95%"
    s1.notes_slide.notes_text_frame.text = "발표자 노트"

    s2 = prs.slides.add_slide(blank)
    s2.shapes.add_textbox(Emu(100), Emu(180000), Emu(3000000), Emu(300000)).text_frame.text = "연구 추진 배경"

    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def test_extract_orders_text_top_to_bottom(tmp_path):
    slides = extract(_deck(tmp_path))
    assert [s["no"] for s in slides] == [1, 2]
    assert slides[0]["blocks"][0] == ["주요 연구내용"]
    assert slides[0]["blocks"][1] == ["최적화를 위한 하이브리드 차량 축약 모델 구성", "1) 하이브리드 차량 축약 모델 구성"]


def test_extract_reads_tables_and_notes(tmp_path):
    first = extract(_deck(tmp_path))[0]
    assert first["tables"] == [[["항목", "값"], ["SoC 상한", "95%"]]]
    assert first["notes"] == "발표자 노트"


def test_render_keeps_wording_and_marks_parts(tmp_path):
    out = render(extract(_deck(tmp_path)))
    assert "=== p1" in out
    assert "[1] 주요 연구내용" in out
    assert "    1) 하이브리드 차량 축약 모델 구성" in out
    assert "[표] SoC 상한 | 95%" in out
    assert "[노트] 발표자 노트" in out
    assert "=== p2" in out


def test_parse_pages():
    assert parse_pages("1-3,7") == {1, 2, 3, 7}


def test_main_out_writes_utf8_lf_and_creates_folder(tmp_path, capsys):
    from extract_slides_text import main

    out = tmp_path / "1_extract" / "slides.md"
    assert main([str(_deck(tmp_path)), "--out", str(out)]) == 0
    data = out.read_bytes()
    assert b"\r\n" not in data
    text = data.decode("utf-8")
    assert "[1] 주요 연구내용" in text
    assert capsys.readouterr().out == ""


def test_main_default_writes_stdout(tmp_path, capsys):
    from extract_slides_text import main

    assert main([str(_deck(tmp_path))]) == 0
    assert "=== p1" in capsys.readouterr().out
